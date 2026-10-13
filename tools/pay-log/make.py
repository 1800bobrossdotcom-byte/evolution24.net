#!/usr/bin/env python3
"""Evolution24 pay logs: one simple workbook per person, in the brand's colours, for Google Sheets.

  python3 tools/pay-log/make.py                          a blank template in tools/pay-log/out/
  python3 tools/pay-log/make.py people.json [out dir]    one file per person in people.json, and a template

people.json holds names and rates of pay, so keep it out of the repository; people.example.json
shows its shape.

Three tabs:
  Pay log          the person, their rate of pay and (if they have two) their role; then a block
                   for each pay week (Saturday to Friday) with a row per day, already dated: six
                   property and hours slots across it, the day's hours and pay worked out, and a
                   note; a total for each week, and the grand total at the top
  Weekly summary   hours and pay for each property, week by week, with totals
  Monthly summary  the same, month by month; new properties are added here
"""
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
LOGO = HERE.parent.parent / 'source' / 'logo' / 'sheet-logo.png'

INK, TAN, CREAM, PAPER, ACCENT = '151613', 'DEBB92', 'E5E6D3', 'F4F1EA', '8A6232'   # assets/css/main.css
INPUT, LINE, MUTED, WHITE, FONT = 'F8EEDF', 'D8D3C6', '6B6A63', 'FFFFFF', 'Arial'
PROPERTIES = ['301 Central Avenue', '145 South Fitzhugh', 'Water Street', 'The Biltmore', '181 St. Paul',
              'Charlotte Square', '2215 James Street', '121 Park Drive', '379 South Main', '561 South Main',
              '31 Genesee Street', '440 Armstrong', '444 Armstrong', '452 Armstrong', '4940 Hillcrest',
              '300 Alexander', 'Office']
SLOTS = 20                     # property rows on Monthly summary: the list above, and room to add more
LOG, WEEK, MONTH = 'Pay log', 'Weekly summary', 'Monthly summary'
WEEKS, MONTHS = 52, 12
PAIRS = 6                      # property and hours slots in each day's row
FIRST = 11                     # Pay log: the first row of the first week
TOTAL = 'Week total'
NAMES = f"'{MONTH}'!$A$9:$A${8 + SLOTS}"   # the property list

MONEY = '$#,##0.00;-$#,##0.00;"–"'
HOURS = '0.0#;-0.0#;"–"'
THIN = Side(style='thin', color=LINE)
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# The Pay log's columns: A the date; then a property and its hours for each slot (B C, D E, ... L M);
# the day's hours, pay and note; and, hidden, the pay per hour that day and each slot's pay.
PROP = [get_column_letter(2 + 2 * k) for k in range(PAIRS)]
HRS = [get_column_letter(3 + 2 * k) for k in range(PAIRS)]
DAY_HOURS, DAY_PAY, NOTE, PER_HOUR = (get_column_letter(2 + 2 * PAIRS + i) for i in range(4))
SLOT_PAY = [get_column_letter(6 + 2 * PAIRS + k) for k in range(PAIRS)]
LAST_COL = SLOT_PAY[-1]


def fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


def font(size=10, color=INK, bold=False, italic=False):
    return Font(name=FONT, size=size, color=color, bold=bold, italic=italic)


def put(ws, cell, value=None, f=None, bg=None, fmt=None, align=None, border=None):
    c = ws[cell]
    if value is not None:
        c.value = value
    c.font = f or font()
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def validate(ws, cells, **kw):
    dv = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    for c in cells:
        dv.add(c)
    ws.add_data_validation(dv)


def band(ws, title, last_col):
    """The ink band: the logo, then the title."""
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 45, 22, 4
    for c in range(1, last_col + 1):
        letter = get_column_letter(c)
        ws[f'{letter}1'].fill = ws[f'{letter}2'].fill = fill(INK)
        ws[f'{letter}3'].fill = fill(TAN)
    put(ws, 'A2', title, font(12, CREAM, True), INK, align=Alignment(vertical='center', indent=1))
    if LOGO.exists():
        img = Image(str(LOGO))
        w = 176
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=pixels_to_EMU(12), row=0, rowOff=pixels_to_EMU(11)),
                                   ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(round(w * img.height / img.width))))
        ws.add_image(img)


def heading_row(ws, row, labels, height=24):
    ws.row_dimensions[row].height = height
    for letter, text, fmt in labels:
        put(ws, f'{letter}{row}', text, font(9, CREAM, True), INK, fmt,
            Alignment(horizontal='center', vertical='center', wrap_text=True))


def first_saturday(person):
    start = dt.date.fromisoformat(person['log_starts']) if person.get('log_starts') else dt.date.today()
    return start - dt.timedelta(days=(start.weekday() - 5) % 7)


def layout(saturday):
    """Where everything sits on the Pay log: per week, its heading, its day rows and its total."""
    weeks, r, month = [], FIRST, None
    for w in range(WEEKS):
        sat = saturday + dt.timedelta(days=7 * w)
        fri = sat + dt.timedelta(days=6)
        week = {'sat': sat, 'fri': fri, 'month_row': None}
        if (fri.year, fri.month) != month:
            month = (fri.year, fri.month)
            week['month_row'] = r
            r += 1
        week['heading'] = r
        week['days'] = [(r + 1 + d, sat + dt.timedelta(days=d)) for d in range(7)]
        week['total'] = r + 8
        weeks.append(week)
        r += 10   # heading, seven days, the total, and a gap
    return weeks


def pay_log(ws, person, weeks):
    shown = 2 + 2 * PAIRS + 2   # A to the note
    band(ws, '="Pay log  ·  "&B5&IF(B7="","","  ·  "&B7)', shown)
    ws.column_dimensions['A'].width = 13
    for k in range(PAIRS):
        ws.column_dimensions[PROP[k]].width = 17
        ws.column_dimensions[HRS[k]].width = 6.5
    ws.column_dimensions[DAY_HOURS].width = 9
    ws.column_dimensions[DAY_PAY].width = 11
    ws.column_dimensions[NOTE].width = 34
    for c in [PER_HOUR] + SLOT_PAY:
        ws.column_dimensions[c].hidden = True

    put(ws, 'A5', 'Person', align=Alignment(vertical='center'))
    ws.merge_cells('B5:D5')
    put(ws, 'B5', person.get('name') or None, font(11, bold=True), INPUT, align=Alignment(vertical='center', indent=1),
        border=GRID)
    put(ws, 'A6', 'Rate of pay', align=Alignment(vertical='center'))
    ws.merge_cells('B6:C6')
    put(ws, 'B6', person.get('rate'), font(11, bold=True), INPUT, MONEY, Alignment(vertical='center'), GRID)
    put(ws, 'D6', 'per week' if person.get('per') == 'week' else 'per hour', font(bold=True), INPUT,
        align=Alignment(horizontal='center', vertical='center'), border=GRID)
    validate(ws, ['D6'], type='list', formula1='"per hour,per week"')
    validate(ws, ['B6'], type='decimal', operator='greaterThanOrEqual', formula1='0', error='An amount in dollars.')
    put(ws, 'A7', 'Role', align=Alignment(vertical='center'))
    ws.merge_cells('B7:D7')
    put(ws, 'B7', person.get('role') or None, font(), INPUT, align=Alignment(vertical='center', indent=1), border=GRID)
    for c in 'CD':
        for r in (5, 7):
            ws[f'{c}{r}'].fill, ws[f'{c}{r}'].border = fill(INPUT), GRID
    ws['C6'].fill, ws['C6'].border = fill(INPUT), GRID
    ws.merge_cells(f'F5:{NOTE}7')
    put(ws, 'F5', 'Fill in your rate of pay. Then, for each day you work, the properties and the hours at each: up to '
                  'six a day, across the row. The day\'s hours and pay work themselves out. On a weekly rate, the week is '
                  'shared across its properties by hours. Role is optional: for anyone with two roles at two rates, who '
                  'keeps a pay log for each.',
        font(9, MUTED, italic=True), align=Alignment(wrap_text=True, vertical='center', indent=1))
    ws.row_dimensions[5].height = ws.row_dimensions[6].height = ws.row_dimensions[7].height = 22

    heading_row(ws, 8, [('A', 'Date', None), (DAY_HOURS, 'Day hours', None), (DAY_PAY, 'Pay', None),
                        (NOTE, 'Work done (optional)', None)] +
                [(PROP[k], f'Property {k + 1}', None) for k in range(PAIRS)] + [(HRS[k], 'Hrs', None) for k in range(PAIRS)])
    whole = lambda c: f'${c}${FIRST}:${c}${weeks[-1]["total"]}'
    for c in range(1, shown + 1):
        put(ws, f'{get_column_letter(c)}9', None, bg=CREAM, border=GRID)
    put(ws, 'A9', 'Total', font(11, bold=True), CREAM, align=Alignment(indent=1, vertical='center'), border=GRID)
    put(ws, f'{DAY_HOURS}9', f'=SUMIF({whole("A")},"{TOTAL}",{whole(DAY_HOURS)})', font(11, bold=True), CREAM, HOURS,
        Alignment(vertical='center'), GRID)
    put(ws, f'{DAY_PAY}9', f'=SUMIF({whole("A")},"{TOTAL}",{whole(DAY_PAY)})', font(11, bold=True), CREAM, MONEY,
        Alignment(vertical='center'), GRID)
    ws.row_dimensions[9].height = 22

    for week in weeks:
        if week['month_row']:
            r = week['month_row']
            ws.row_dimensions[r].height = 26
            put(ws, f'A{r}', week['fri'].strftime('%B %Y'), font(12, ACCENT, True), align=Alignment(vertical='bottom'))
        h, t = week['heading'], week['total']
        for c in range(1, shown + 1):
            put(ws, f'{get_column_letter(c)}{h}', None, bg=PAPER)
        put(ws, f'A{h}', f"Week of {week['sat']:%a %-m/%-d} – {week['fri']:%a %-m/%-d/%Y}", font(bold=True), PAPER,
            align=Alignment(indent=1, vertical='center'))
        ws.row_dimensions[h].height = 20
        first, last = week['days'][0][0], week['days'][-1][0]
        for r, day in week['days']:
            put(ws, f'A{r}', day, font(), WHITE, 'ddd m/d', Alignment(indent=1), GRID)
            for k in range(PAIRS):
                put(ws, f'{PROP[k]}{r}', None, font(), INPUT, align=Alignment(indent=1), border=GRID)
                put(ws, f'{HRS[k]}{r}', None, font(), INPUT, HOURS, border=GRID)
                put(ws, f'{SLOT_PAY[k]}{r}', f'=IF(OR({HRS[k]}{r}="",{PER_HOUR}{r}=""),"",{HRS[k]}{r}*{PER_HOUR}{r})',
                    font(color=MUTED), fmt=MONEY)
            put(ws, f'{DAY_HOURS}{r}', '=SUM(' + ','.join(f'{c}{r}' for c in HRS) + ')', font(), WHITE, HOURS, border=GRID)
            put(ws, f'{DAY_PAY}{r}', f'=IF(OR({DAY_HOURS}{r}=0,{PER_HOUR}{r}=""),"",{DAY_HOURS}{r}*{PER_HOUR}{r})',
                font(), WHITE, MONEY, border=GRID)
            put(ws, f'{NOTE}{r}', None, font(), INPUT, align=Alignment(indent=1), border=GRID)
            # Pay per hour this day: the rate, or on a weekly rate the week's pay over the week's hours.
            put(ws, f'{PER_HOUR}{r}', f'=IF($B$6="","",IF($D$6="per week",IF({DAY_HOURS}${t}=0,"",$B$6/{DAY_HOURS}${t}),$B$6))',
                font(color=MUTED), fmt=MONEY)
        for c in range(1, shown + 1):
            put(ws, f'{get_column_letter(c)}{t}', None, bg=CREAM, border=GRID)
        put(ws, f'A{t}', TOTAL, font(bold=True), CREAM, align=Alignment(indent=1), border=GRID)
        put(ws, f'{DAY_HOURS}{t}', f'=SUM({DAY_HOURS}{first}:{DAY_HOURS}{last})', font(bold=True), CREAM, HOURS, border=GRID)
        put(ws, f'{DAY_PAY}{t}', f'=SUM({DAY_PAY}{first}:{DAY_PAY}{last})', font(bold=True), CREAM, MONEY, border=GRID)

    days = [r for week in weeks for r, _ in week['days']]
    validate(ws, [f'{PROP[k]}{r}' for r in days for k in range(PAIRS)], type='list', formula1=NAMES, errorStyle='warning',
             errorTitle='Property', error='Not on the property list (add new ones on Monthly summary). Keep it anyway?')
    validate(ws, [f'{HRS[k]}{r}' for r in days for k in range(PAIRS)], type='decimal', operator='between', formula1='0',
             formula2='24', errorTitle='Hours', error='Hours as a number: 7.5 is seven and a half.')
    ws.freeze_panes = 'B10'
    ws.print_title_rows = '8:8'
    ws.page_setup.orientation = 'landscape'


LOG_COL = lambda c, last: f"'{LOG}'!${c}${FIRST}:${c}${last}"


def by_property(weeks, what, prop, span):
    """Hours or pay at one property: the sum over the six slots of each day."""
    last = weeks[-1]['total']
    values = HRS if what == 'hours' else SLOT_PAY
    return '+'.join(f'SUMIFS({LOG_COL(v, last)},{LOG_COL(p, last)},{prop},{span})' for v, p in zip(values, PROP))


def weekly(ws, weeks):
    last = weeks[-1]['total']
    dates = LOG_COL('A', last)
    cols = [get_column_letter(2 + i) for i in range(SLOTS)]
    other, total = get_column_letter(2 + SLOTS), get_column_letter(3 + SLOTS)
    band(ws, f"=\"Weekly summary  ·  \"&'{LOG}'!B5", 3 + SLOTS)
    ws.column_dimensions['A'].width = 16
    for c in cols + [other]:
        ws.column_dimensions[c].width = 10
    ws.column_dimensions[total].width = 12
    put(ws, 'A5', 'Each pay week runs Saturday to Friday. Properties come from the list on Monthly summary.',
        font(9, MUTED, italic=True))

    def grid(top, heading, what, fmt):
        day_total = LOG_COL(DAY_HOURS if what == 'hours' else DAY_PAY, last)
        put(ws, f'A{top}', heading, font(12, ACCENT, True))
        h = top + 1
        heading_row(ws, h, [('A', 'Week ending', None), (other, 'Other or none', None), (total, 'Total', None)] +
                    [(c, f"=IF('{MONTH}'!A{9 + i}=\"\",\"\",'{MONTH}'!A{9 + i})", None) for i, c in enumerate(cols)], 40)
        for i, week in enumerate(weeks):
            r = h + 1 + i
            span = f'{dates},">="&($A{r}-6),{dates},"<="&$A{r}'
            put(ws, f'A{r}', week['fri'], font(), WHITE, 'ddd m/d/yy', Alignment(indent=1), GRID)
            for c in cols:
                put(ws, f'{c}{r}', f'=IF({c}${h}="","",{by_property(weeks, what, f"{c}${h}", span)})', font(), WHITE, fmt,
                    border=GRID)
            put(ws, f'{total}{r}', f'=SUMIFS({day_total},{span})', font(bold=True), CREAM, fmt, border=GRID)
            put(ws, f'{other}{r}', f'={total}{r}-SUM({cols[0]}{r}:{cols[-1]}{r})', font(color=MUTED, italic=True), WHITE, fmt,
                border=GRID)
        return h + WEEKS

    end = grid(7, 'Hours', 'hours', HOURS)
    grid(end + 2, 'Pay', 'pay', MONEY)
    ws.freeze_panes = 'B9'


def monthly(ws, weeks):
    last = weeks[-1]['total']
    dates = LOG_COL('A', last)
    band(ws, f"=\"Monthly summary  ·  \"&'{LOG}'!B5", 2 + MONTHS)
    ws.column_dimensions['A'].width = 26
    for c in range(2, 2 + MONTHS):
        ws.column_dimensions[get_column_letter(c)].width = 10.5
    end_col = get_column_letter(2 + MONTHS)
    ws.column_dimensions[end_col].width = 12
    put(ws, 'A5', 'To add a property, type it in an empty tan row under Hours; it joins the list on the Pay log.',
        font(9, MUTED, italic=True))
    sat = weeks[0]['sat']
    months = [dt.date(sat.year + (sat.month - 1 + m) // 12, (sat.month - 1 + m) % 12 + 1, 1) for m in range(MONTHS)]

    def grid(top, heading, what, fmt, names_from=None):
        day_total = LOG_COL(DAY_HOURS if what == 'hours' else DAY_PAY, last)
        put(ws, f'A{top}', heading, font(12, ACCENT, True))
        h = top + 1
        heading_row(ws, h, [('A', 'Property', None), (end_col, 'Total', None)] +
                    [(get_column_letter(2 + m), month, 'mmm yyyy') for m, month in enumerate(months)])
        rows = range(h + 1, h + 1 + SLOTS)
        other, total = h + 1 + SLOTS, h + 2 + SLOTS
        for i, r in enumerate(rows):
            if names_from is None:
                put(ws, f'A{r}', PROPERTIES[i] if i < len(PROPERTIES) else None, font(), INPUT, align=Alignment(indent=1),
                    border=GRID)
            else:
                put(ws, f'A{r}', f'=IF(A{names_from + i}="","",A{names_from + i})', font(), WHITE, align=Alignment(indent=1),
                    border=GRID)
        put(ws, f'A{other}', 'Other or no property', font(color=MUTED, italic=True), WHITE, align=Alignment(indent=1), border=GRID)
        put(ws, f'A{total}', 'Total', font(bold=True), CREAM, align=Alignment(indent=1), border=GRID)
        for m in range(MONTHS):
            c = get_column_letter(2 + m)
            span = f'{dates},">="&{c}${h},{dates},"<"&EDATE({c}${h},1)'
            for r in rows:
                put(ws, f'{c}{r}', f'=IF($A{r}="","",{by_property(weeks, what, f"$A{r}", span)})', font(), WHITE, fmt,
                    border=GRID)
            put(ws, f'{c}{total}', f'=SUMIFS({day_total},{span})', font(bold=True), CREAM, fmt, border=GRID)
            put(ws, f'{c}{other}', f'={c}{total}-SUM({c}{rows[0]}:{c}{rows[-1]})', font(color=MUTED, italic=True), WHITE, fmt,
                border=GRID)
        for r in list(rows) + [other, total]:
            put(ws, f'{end_col}{r}', f'=IF($A{r}="","",SUM(B{r}:{get_column_letter(1 + MONTHS)}{r}))', font(bold=True),
                CREAM if r == total else WHITE, fmt, border=GRID)
        return rows[0], total

    names, end = grid(7, 'Hours', 'hours', HOURS)
    grid(end + 2, 'Pay', 'pay', MONEY, names_from=names)
    ws.freeze_panes = 'B9'


def build(person, path):
    weeks = layout(first_saturday(person))
    wb = Workbook()
    wb._fonts[0] = Font(name=FONT, size=10)   # the default for cells nobody styled
    wb.active.title = LOG
    pay_log(wb.active, person, weeks)
    weekly(wb.create_sheet(WEEK), weeks)
    monthly(wb.create_sheet(MONTH), weeks)
    for ws in wb.worksheets:
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb[WEEK].page_setup.orientation = wb[MONTH].page_setup.orientation = 'landscape'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def file_name(person):
    name = (person.get('name') or 'Template').rstrip('.')
    role = person.get('role')
    return f"Evolution24 Pay Log - {name} ({role}).xlsx" if role else f"Evolution24 Pay Log - {name}.xlsx"


def main(argv):
    if len(argv) > 1:
        data = json.loads(Path(argv[1]).read_text())
        out = Path(argv[2]) if len(argv) > 2 else HERE / 'out'
        people = [{'log_starts': data.get('log_starts'), **p} for p in data['people']]
        people.append({'log_starts': data.get('log_starts')})
    else:
        out, people = HERE / 'out', [{}]
    for p in people:
        print(build(p, out / file_name(p)))


if __name__ == '__main__':
    main(sys.argv)
