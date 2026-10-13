#!/usr/bin/env python3
"""Fills two pay logs (one hourly, one on a weekly rate) with made-up lines, lets LibreOffice
calculate them, and checks every line's pay, each week's total, the grand total, and the weekly
and monthly summaries against the same sums worked out here in Python.

  python3 tools/pay-log/test.py
  python3 tools/pay-log/test.py --fixture out.json   filled pay logs as JSON, for scripts/test-pay-report.cjs

Needs LibreOffice Calc (soffice) on the PATH.
"""
import datetime as dt
import json
import os
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make  # noqa: E402

D = dt.date
ENTRIES = [  # date, property, hours
    (D(2026, 10, 10), 'Water Street', 3),
    (D(2026, 10, 12), 'Water Street', 5.5),
    (D(2026, 10, 12), '145 South Fitzhugh', 2.5),     # two properties on one day
    (D(2026, 10, 12), 'Office / all properties', 0.75),                # and a third
    (D(2026, 10, 13), 'The Biltmore', 8),
    (D(2026, 10, 15), '301 Central Avenue', 1),      # all six slots on one day
    (D(2026, 10, 15), '181 St. Paul', 1),
    (D(2026, 10, 15), '2215 James Street', 0.5),
    (D(2026, 10, 15), '121 Park Drive', 0.5),
    (D(2026, 10, 15), '31 Genesee Street', 1),
    (D(2026, 10, 15), '561 South Main', 1.5),
    (D(2026, 10, 16), 'Office / all properties', 1.25),
    (D(2026, 10, 19), 'Charlotte Square', 9),
    (D(2026, 10, 31), 'Water Street', 4),             # a Saturday: the next pay week, still October
    (D(2026, 11, 2), 'Water Street', 6),
    (D(2026, 11, 3), 'A garage on Main', 2),          # not on the list: counted under Other
    (D(2026, 11, 4), '', 1),                           # no property: also Other
    (D(2027, 9, 30), '379 South Main', 7),            # the twelfth month
]
PEOPLE = [{'name': 'Hourly Test', 'rate': 21.5, 'per': 'hour', 'log_starts': '2026-10-12'},   # a Monday: back to Sat 10/10
          {'name': 'Weekly Test', 'rate': 1000, 'per': 'week', 'log_starts': '2026-10-10'}]
OTHER = 'Other or no property'


def week_end(d):
    return d + dt.timedelta(days=(4 - d.weekday()) % 7)   # the Friday on or after d


def expected(person):
    week_hours = defaultdict(float)
    for d, _, h in ENTRIES:
        week_hours[week_end(d)] += h
    pay = [person['rate'] * h if person['per'] == 'hour' else person['rate'] * h / week_hours[week_end(d)]
           for d, _, h in ENTRIES]
    by_week, by_month = defaultdict(float), defaultdict(float)   # (what, period, property) -> sum
    for (d, p, h), money in zip(ENTRIES, pay):
        prop = p if p in make.PROPERTIES else OTHER
        for what, value in (('hours', h), ('pay', money)):
            for key in (prop, 'Total'):
                by_week[what, week_end(d), key] += value
                by_month[what, (d.year, d.month), key] += value
    return pay, by_week, by_month


def recalc(path):
    """Calculate every formula with LibreOffice and save the values into the file."""
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / 'profile'
        env = {**os.environ, 'SAL_USE_VCLPLUGIN': 'svp'}
        run = lambda *a: subprocess.run(['soffice', '--headless', f'-env:UserInstallation={profile.as_uri()}', *a],
                                        env=env, capture_output=True, timeout=180)
        run('--terminate_after_init')
        (profile / 'user' / 'basic' / 'Standard' / 'Module1.xba').write_text(
            '<?xml version="1.0" encoding="UTF-8"?>\n<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">\n'
            '<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">\n'
            'Sub Recalc\n  ThisComponent.calculateAll()\n  ThisComponent.store()\n  ThisComponent.close(True)\nEnd Sub\n</script:module>')
        run('macro:///Standard.Module1.Recalc', str(path))


def filled(person, entries, path):
    """A pay log for `person` with `entries` in its day rows' slots, calculated. Returns the layout and the slots used."""
    weeks = make.layout(make.first_saturday(person))
    free = defaultdict(list)                  # each day's slots, in order: (row, slot)
    for week in weeks:
        for r, day in week['days']:
            free[day] += [(r, k) for k in range(make.PAIRS)]
    placed = [free[d].pop(0) for d, _, _ in entries]
    make.build(person, path)
    wb = load_workbook(path)
    ws = wb[make.LOG]
    for (r, k), (d, p, h) in zip(placed, entries):
        ws[f'{make.PROP[k]}{r}'], ws[f'{make.HRS[k]}{r}'] = p or None, h
        if p and p.startswith('='):
            ws[f'{make.PROP[k]}{r}'].data_type = 's'   # text that only looks like a formula
    wb.save(path)
    recalc(path)
    return weeks, placed


def check(person, keep=None):
    problems = []

    def same(label, value, want):
        if abs((value or 0) - want) > 0.005:
            problems.append(f"{person['name']}: {label}: sheet {value!r}, expected {want!r}")

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / 'test.xlsx'
        weeks, placed = filled(person, ENTRIES, path)
        if keep:
            Path(keep).write_bytes(path.read_bytes())
        got = load_workbook(path, data_only=True)
    if weeks[0]['sat'] != D(2026, 10, 10):
        problems.append(f"first week starts {weeks[0]['sat']}, expected Sat 10/10/26")

    errors = [f'{s.title}!{c.coordinate}' for s in got.worksheets for row in s.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith('#')]
    problems += [f'formula error at {e}' for e in errors[:10]]
    pay, by_week, by_month = expected(person)
    log, wk, mo = got[make.LOG], got[make.WEEK], got[make.MONTH]

    day_pay = defaultdict(float)
    for (r, k), money in zip(placed, pay):
        same(f'row {r} slot {k + 1} pay', log[f'{make.SLOT_PAY[k]}{r}'].value, money)
        day_pay[r] += money
    for r, money in day_pay.items():
        same(f'row {r} day pay', log[f'{make.DAY_PAY}{r}'].value, money)
    for week in weeks:
        t = week['total']
        same(f"week {week['fri']} total hours", log[f'{make.DAY_HOURS}{t}'].value, by_week['hours', week['fri'], 'Total'])
        same(f"week {week['fri']} total pay", log[f'{make.DAY_PAY}{t}'].value, by_week['pay', week['fri'], 'Total'])
        for r, day in week['days']:
            if log[f'A{r}'].value.date() != day:
                problems.append(f'row {r} is dated {log[f"A{r}"].value}, expected {day}')
    same('grand total hours', log[f'{make.DAY_HOURS}9'].value, sum(h for _, _, h in ENTRIES))
    same('grand total pay', log[f'{make.DAY_PAY}9'].value, sum(pay))

    for h, what in ((8, 'hours'), (63, 'pay')):
        names = {c: wk.cell(h, c).value for c in range(2, 4 + make.SLOTS)}
        names[2 + make.SLOTS] = OTHER
        for i, week in enumerate(weeks):
            r = h + 1 + i
            if wk.cell(r, 1).value.date() != week['fri']:
                problems.append(f'weekly summary row {r} is {wk.cell(r, 1).value}, expected {week["fri"]}')
            for c, name in names.items():
                if name:
                    same(f"weekly {what} {week['fri']} {name}", wk.cell(r, c).value, by_week[what, week['fri'], name])
    for h, what in ((8, 'hours'), (33, 'pay')):
        months = [(mo.cell(h, c).value.year, mo.cell(h, c).value.month) for c in range(2, 2 + make.MONTHS)]
        for r in range(h + 1, h + 3 + make.SLOTS):
            name = mo.cell(r, 1).value
            if not name:
                continue
            for c, month in enumerate(months, start=2):
                same(f'monthly {what} {name} {month}', mo.cell(r, c).value, by_month[what, month, name])
            same(f'monthly {what} {name} total', mo.cell(r, 2 + make.MONTHS).value,
                 sum(v for (w, m, n), v in by_month.items() if w == what and n == name and m in months))
    return problems


# For the labor report's test: three people, as their pay logs read once calculated.
FIXTURE = [
    ({'name': 'Alex Test', 'role': 'Maintenance', 'rate': 21.5, 'per': 'hour', 'log_starts': '2026-10-10'}, ENTRIES),
    ({'name': 'Alex Test', 'role': 'Management', 'rate': 38, 'per': 'hour', 'log_starts': '2026-10-10'}, [   # a second role
        (D(2026, 10, 14), 'Office / all properties', 3), (D(2026, 10, 21), 'Water Street', 2)]),
    ({'name': 'Sam Test', 'rate': 1000, 'per': 'week', 'log_starts': '2026-10-10'}, [
        (D(2026, 10, 12), 'Water Street', 8), (D(2026, 10, 13), 'Charlotte Square', 6),
        (D(2026, 10, 13), 'Water Street', 2), (D(2026, 10, 20), 'The Biltmore', 9), (D(2026, 11, 9), 'Office', 4)]),   # the old name: still shared by the report
    ({'name': 'Jo Test', 'per': 'hour', 'log_starts': '2026-10-10'}, [   # no rate of pay set
        (D(2026, 10, 14), 'Water Street', 3), (D(2026, 10, 15), '=HYPERLINK("http://example.com","x")', 1)]),
]


def fixture(out):
    def plain(v):
        if isinstance(v, dt.datetime):
            return {'date': v.date().isoformat()}
        return v
    logs = []
    with tempfile.TemporaryDirectory() as tmp:
        for person, entries in FIXTURE:
            path = Path(tmp) / make.file_name(person)
            filled(person, entries, path)
            ws = load_workbook(path, data_only=True)[make.LOG]
            width = ws[f'{make.LAST_COL}1'].column
            rows = [[plain(c.value) for c in row] for row in ws.iter_rows(min_row=1, max_col=width)]
            logs.append({'name': path.stem, 'person': person['name'], 'role': person.get('role'), 'rows': rows,
                         'entries': [[d.isoformat(), p, h] for d, p, h in entries],
                         'rate': person.get('rate'), 'per': person.get('per')})
    Path(out).write_text(json.dumps(logs))
    print(f'{out}: {len(logs)} pay logs')


def main():
    if sys.argv[1:2] == ['--fixture']:
        return fixture(sys.argv[2])
    problems = []
    for i, person in enumerate(PEOPLE):
        problems += check(person, keep=os.environ.get('KEEP') and f"{os.environ['KEEP']}-{i}.xlsx")
    if problems:
        print('\n'.join(problems[:40]))
        sys.exit(f'{len(problems)} problems')
    print(f'pay log: hourly and weekly rates, {len(ENTRIES)} lines: every line, week, month and total matches')


if __name__ == '__main__':
    main()
